from __future__ import annotations

import logging
from pathlib import Path
from types import TracebackType
from typing import Any, cast

import pandas as pd
from pytest import MonkeyPatch

import main as app_main


def make_valid_record(record_id: int) -> app_main.RowData:
    return {
        "ID": record_id,
        "TIPO": "ÚNICA",
        "DIMENSÃO": "1. Ambiental",
        "TEMA": "1.1. Água e Efluentes",
        "TÓPICO": "1.1.1. Água e Efluentes",
        "APLICAÇÃO SETORIAL": "Aplicável",
        "CÓDIGO QUESTÃO": "1.1.1.1",
        "ANO": 2026,
        "PERGUNTA": "Pergunta",
        "SUB-OPÇÃO": None,
        "ORIENTAÇÕES": "Selecione uma alternativa",
        "TIPO DE EVIDÊNCIA": "Apenas Evidência Pública",
        "EXEMPLOS DE EVIDÊNCIA": "Relatório anual",
        "STATUS": None,
        "PONTO FOCAL": None,
        "RESPONSÁVEL": None,
        "ENVIAR EVIDÊNCIA": None,
        "PTS. CABÍVEIS": None,
        "DISTRIBUIÇÃO": None,
        "RESPOSTA": None,
        "EVIDÊNCIA EMPRESA": None,
        "PTS. OBTIDOS": None,
        "%": None,
        "GRI/ISO": None,
        "TRIAL": None,
    }


def test_configure_logging_creates_handlers_and_log_file() -> None:
    logger = app_main.configure_logging()

    assert logger.name == app_main.LOGGER_NAME
    assert len(logger.handlers) == 2
    assert (app_main.PROJECT_ROOT / "logs" / "app.log").exists()


def test_extract_lines_from_pdf_reads_and_strips_lines(monkeypatch: MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, content: str) -> None:
            self._content = content

        def extract_text(self, layout: bool = False) -> str:
            _ = layout
            return self._content

    class FakePDF:
        def __init__(self) -> None:
            self.pages = [FakePage(" linha 1\n\nlinha 2 "), FakePage("linha 3")]

        def __enter__(self) -> "FakePDF":
            return self

        def __exit__(
            self,
            exc_type: type[BaseException] | None,
            exc: BaseException | None,
            tb: TracebackType | None,
        ) -> bool:
            return False

    class FakePDFPlumber:
        def open(self, _path: Path) -> FakePDF:
            return FakePDF()

    monkeypatch.setattr(app_main, "pdfplumber", FakePDFPlumber())

    lines = app_main.extract_lines_from_pdf(Path("fake.pdf"))

    assert lines == ["linha 1", "linha 2", "linha 3"]


def test_build_dataframe_orders_columns_and_fills_missing() -> None:
    frame = app_main.build_dataframe([make_valid_record(1)])

    assert list(frame.columns) == app_main.COLUMN_ORDER
    assert frame.iloc[0]["ID"] == 1


def test_apply_basic_table_format_returns_without_rows(tmp_path: Path) -> None:
    output_path = tmp_path / "empty.xlsx"
    cast(Any, pd.DataFrame()).to_excel(output_path, index=False)

    app_main.apply_basic_table_format(output_path, row_count=0, column_count=0)

    assert output_path.exists()


def test_apply_basic_table_format_applies_table_and_freeze_panes(tmp_path: Path) -> None:
    output_path = tmp_path / "full.xlsx"
    cast(Any, pd.DataFrame([{"A": 1, "B": 2}, {"A": 3, "B": 4}])).to_excel(output_path, index=False)

    app_main.apply_basic_table_format(output_path, row_count=2, column_count=2)

    from openpyxl import load_workbook

    worksheet = load_workbook(output_path).active
    assert worksheet is not None
    assert worksheet.freeze_panes == "A2"
    assert len(worksheet.tables) == 1


def test_main_returns_1_when_cli_args_invalid(monkeypatch: MonkeyPatch) -> None:
    def fake_parse_cli_args(argv: list[str], logger: logging.Logger) -> None:
        _ = (argv, logger)
        return None

    monkeypatch.setattr(app_main, "parse_cli_args", fake_parse_cli_args)

    assert app_main.main() == 1


def test_main_success_flow_returns_0(monkeypatch: MonkeyPatch, tmp_path: Path) -> None:
    output_path = tmp_path / "saida.xlsx"
    sample_records = [make_valid_record(1), make_valid_record(2)]

    def fake_parse_cli_args(argv: list[str], logger: logging.Logger) -> tuple[list[Path], Path]:
        _ = (argv, logger)
        return [Path("a.pdf")], output_path

    def fake_extract_lines_from_pdf(_path: Path) -> list[str]:
        return ["linha"]

    def fake_parse_pdf_lines(
        lines: list[str],
        input_path: Path | None,
        logger: logging.Logger | None,
    ) -> list[app_main.RowData]:
        _ = (lines, input_path, logger)
        return sample_records

    monkeypatch.setattr(app_main, "parse_cli_args", fake_parse_cli_args)
    monkeypatch.setattr(app_main, "extract_lines_from_pdf", fake_extract_lines_from_pdf)
    monkeypatch.setattr(app_main, "parse_pdf_lines", fake_parse_pdf_lines)

    assert app_main.main() == 0
    assert output_path.exists()


def test_main_generates_alternative_output_when_file_exists(
    monkeypatch: MonkeyPatch,
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "saida.xlsx"
    cast(Any, pd.DataFrame([{"A": 1}])).to_excel(output_path, index=False)
    sample_records = [make_valid_record(1)]

    def fake_parse_cli_args(argv: list[str], logger: logging.Logger) -> tuple[list[Path], Path]:
        _ = (argv, logger)
        return [Path("a.pdf")], output_path

    def fake_extract_lines_from_pdf(_path: Path) -> list[str]:
        return ["linha"]

    def fake_parse_pdf_lines(
        lines: list[str],
        input_path: Path | None,
        logger: logging.Logger | None,
    ) -> list[app_main.RowData]:
        _ = (lines, input_path, logger)
        return sample_records

    monkeypatch.setattr(app_main, "parse_cli_args", fake_parse_cli_args)
    monkeypatch.setattr(app_main, "extract_lines_from_pdf", fake_extract_lines_from_pdf)
    monkeypatch.setattr(app_main, "parse_pdf_lines", fake_parse_pdf_lines)

    assert app_main.main() == 0
    assert output_path.exists()
    assert (tmp_path / "saida_01.xlsx").exists()


def test_main_returns_1_on_processing_error(monkeypatch: MonkeyPatch, tmp_path: Path) -> None:
    output_path = tmp_path / "saida.xlsx"

    def fake_parse_cli_args(argv: list[str], logger: logging.Logger) -> tuple[list[Path], Path]:
        _ = (argv, logger)
        return [Path("a.pdf")], output_path

    def fake_extract_lines_from_pdf(_path: Path) -> list[str]:
        raise OSError("falha")

    monkeypatch.setattr(app_main, "parse_cli_args", fake_parse_cli_args)
    monkeypatch.setattr(app_main, "extract_lines_from_pdf", fake_extract_lines_from_pdf)

    assert app_main.main() == 1


def test_parse_cli_args_handles_too_few_arguments() -> None:
    logger = logging.getLogger("test-cli-few")

    result = app_main.parse_cli_args(["main.py", "somente_um_arg"], logger)

    assert result is None
